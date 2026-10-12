#!/usr/bin/env python3
"""Build the two-page KICS short paper from the KICS Word template.

The KICS sample is a legacy ``.doc``.  Convert it once to ``.docx`` *and* upgrade
it out of Word 2003 compatibility mode (Word: Save As .docx, reopen, File > Info >
Convert; or COM ``SetCompatibilityMode(15)``) and pass it as ``--template``.
Compatibility mode 11 cannot hold native equations, so Word breaks them on save;
the build refuses such a template.  ``output/kics/kics_template_word2013.docx``
is the converted template.
Paragraph, run, section and column formatting are cloned from the template;
only the text is replaced.  Equations are native Word equations (OMML),
converted LaTeX -> MathML (``latex2mathml``) -> OMML (Office ``MML2OMML.XSL``).

Extra dependencies (not in requirements.txt): python-docx, latex2mathml.
All numbers come from the frozen result files listed in the paper text; see
``docs/review_ai_feedback_2026-09-23.md`` for their provenance.
"""
from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

import docx
import latex2mathml.converter
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Template paragraph indices (KICS sample converted to .docx).
T_TITLE1, T_TITLE2, T_AUTH, T_AFFIL, T_MAIL = 1, 2, 3, 4, 5
T_ETITLE1, T_ETITLE2, T_EAUTH, T_EAFFIL = 7, 8, 9, 10
T_ABS_HEAD, T_ABS = 12, 14
T_SEC_END = 19            # last single-column paragraph (holds the sectPr)
T_HEAD, T_BODY, T_GAP = 20, 21, 22
T_REF_HEAD, T_REF = 32, 33

COLUMN_TWIPS = 4749       # (11909 - 1134 - 851 - 425) / 2

# The KICS sample names 한양신명조, a font that exists only inside Hangul (HWP) and
# is not a Windows font, so Word silently substitutes 바탕.  HY신명조 is the
# TrueType font of the same Hanyang family installed with Hancom Office.
# HY신명조 has poor Latin glyphs (wide hyphen, no "ú"), so Latin text, digits and
# symbols use Times New Roman.
TEMPLATE_FONT = "한양신명조"
BODY_FONT = "HY신명조"
LATIN_FONT = "Times New Roman"


# ---------------------------------------------------------------------------
# text content
# ---------------------------------------------------------------------------
TITLE_KO = ("유한 블록길이 과부하 MU-MISO 하향링크에서", "AoI 최소화를 위한 PPO 기반 RSMA 자원 할당")
TITLE_EN = ("PPO-Based RSMA Resource Allocation for AoI Minimization",
            "in Finite-Blocklength Overloaded MU-MISO Downlink")
AUTHORS_KO = "박장호, 곽도현"
AFFIL_KO = "충북대학교 정보통신공학부, 충북대학교 정보통신공학부"
EMAIL = "2021039085@cbnu.ac.kr, 202103965@cbnu.ac.kr"
AUTHORS_EN = "Park Jang Ho, Gwak Do Hyun"
AFFIL_EN = "Chungbuk Univ., Chungbuk Univ."

ABSTRACT = (
    "본 논문은 사용자 3명, 송신 안테나 2개의 과부하 MU-MISO 하향링크에서 유한 블록길이(FBL) RSMA의 "
    "전력과 common payload 분할을 PPO로 학습하여 장기 평균 AoI를 낮춘다. 목표 조건(SNR 15 dB, blocklength "
    "128)에서 동일한 objective-only 사전학습과 PPO 예산을 적용한 5개 학습 seed의 평균 AoI는 1.463으로, "
    "common stream 없는 SDMA 정책, 검증으로 조정한 고정 규칙 RSMA, 최대 AoI 단일 사용자 전송보다 각각 "
    "16.01%, 11.63%, 26.87% 낮았다. "
    "학습된 정책은 AoI가 큰 사용자에게 private 전력을 집중하고 AoI 순위별 common payload를 다르게 배분했으며, "
    "SNR과 사용자 각도에 따라 "
    "common 전력을 조절했으나, 낮은 SNR과 짧은 blocklength에서는 단일 사용자 전송이 우세했다."
)

# Body blocks: ("h", text) heading, ("p", text) paragraph with $...$ inline math,
# ("eq", latex, number) display equation, ("table", key), ("figure", key), ("gap",).
BODY = [
    ("h", "Ⅰ. 서 론"),
    ("p", "상태 감시·제어 응용에서는 수신 정보의 신선도인 AoI(Age of Information)가 중요한 지표다[1]. 사용자 "
          "수가 송신 안테나보다 많은 과부하 MU-MISO 하향링크에서는 precoding만으로 간섭을 억제하기 어렵고, RSMA는 "
          "메시지 일부를 모든 사용자가 복호하는 common stream으로 보내 간섭 일부를 복호한다[2]. 짧은 패킷에서는 FBL "
          "효과로 성공확률이 SINR과 payload에 따라 급변하고[3] 조건에 따라 common stream을 끄는 편이 유리하며[4], "
          "장기 AoI는 현재 AoI와 시변 채널에 의존한다. 기존 연구[5]는 FBL MU-MISO RSMA에서 DRL로 "
          "precoder·전력·rate split을 공동 학습하였다. 본 논문은 과부하 환경($K=3>N_t=2$)에서 beam 생성 규칙을 "
          "고정하고 SDMA 경계를 포함한 전력·payload 분할만 PPO로 학습해 동일 난수에서 SDMA·AoI 규칙과 비교한다."),
    ("gap",),
    ("h", "Ⅱ. 시스템 모델"),
    ("p", "송신 안테나 $N_t=2$개인 기지국이 단일 안테나 사용자 $K=3$명에게 매 slot 256-bit 상태 update를 "
          "보낸다(generate-at-will). 채널은 ULA 조향 벡터를 LoS 성분으로 하는 Rician 채널($\\kappa=8$ dB)이며 "
          "산란 성분은 상관계수 0.8의 Gauss-Markov 과정을 따른다. 기본 사용자 각도는 "
          "$(-10^{\\circ},0^{\\circ},10^{\\circ})$(각도 폭 20°)이고, 송신기는 산란 성분에 상대 오차 $\\sigma_e=0.2$가 "
          "섞인 추정 채널 $\\hat{\\mathbf{h}}_k$만 안다(전체 채널 NMSE 약 −22.6 dB). 송신 신호는"),
    ("eq", r"\mathbf{x}=\sqrt{\alpha P}\mathbf{u}_c s_c+\sum_{k}\sqrt{(1-\alpha)q_kP}\mathbf{u}_k s_k", "1"),
    ("p", "이다. $P$는 총 전력(SNR $=P/\\sigma^2$), $\\mathbf{u}_k$는 $\\hat{\\mathbf{h}}_k$ 기반 RZF 방향, "
          "$\\mathbf{u}_c$는 추정 채널 행렬의 주 특이벡터다. $P_c=\\alpha P$, $P_k=(1-\\alpha)q_kP$일 때 common "
          "SINR과 SIC 이후 private SINR은"),
    ("eq", r"\gamma_{c,k}=\frac{P_c|\mathbf{h}_k^{\mathrm{H}}\mathbf{u}_c|^2}"
           r"{\sum_{j}P_j|\mathbf{h}_k^{\mathrm{H}}\mathbf{u}_j|^2+\sigma^2}", "2"),
    ("eq", r"\gamma_{p,k}=\frac{P_k|\mathbf{h}_k^{\mathrm{H}}\mathbf{u}_k|^2}"
           r"{\sum_{j\neq k}P_j|\mathbf{h}_k^{\mathrm{H}}\mathbf{u}_j|^2+\sigma^2}", "3"),
    ("p", "이다. 사용자별 $D=256$ bit 중 $\\phi_kD$ bit는 common으로 보내며, $D_c=\\sum_k\\phi_kD$ bit의 "
          "common codeword는 모든 사용자가 복호한다. blocklength $n$에서 $d$ bit의 BLER은 normal approximation[3]"),
    ("eq", r"\epsilon(\gamma,d)=Q\left(\frac{n\log_2(1+\gamma)-d}{\sqrt{nV(\gamma)}}\right)", "4"),
    ("p", "로 계산하고($V(\\gamma)=(1-(1+\\gamma)^{-2})\\log_2^2e$) 성공확률은"),
    ("eq", r"p_k=[1-\epsilon(\gamma_{c,k},D_c)][1-\epsilon(\gamma_{p,k},(1-\phi_k)D)]", "5"),
    ("p", "이다. ACK $A_k(t)\\sim\\mathrm{Bern}(p_k)$가 1이면 AoI $\\Delta_k$는 1로, 아니면 1 증가하며 목표는 "
          "장기 평균 AoI의 최소화다."),
    ("gap",),
    ("h", "Ⅲ. PPO 기반 자원 할당"),
    ("p", "행동 $a_t=(\\alpha,\\{q_k\\},\\{\\phi_k\\})$에서 $\\alpha$와 $\\phi_k$는 Gaussian 잠재변수 $z$를 "
          "$\\mathrm{clip}(1/2+z/6,0,1)$로, $q_k$는 softmax로 변환해 SDMA 경계($\\alpha=\\phi_k=0$)를 정확히 "
          "표현한다. 상태는 정규화 AoI, 이전 ACK와 전력 할당, 추정 채널 기반 이득·상관·beam 이득·nominal SINR, "
          "SNR과 $n$이다. 보상은 $r_t=-\\frac{1}{K}\\sum_k\\Delta_k(t+1)$이며 PPO[6](GAE, 할인율 0.99, clip 0.2, "
          "2×128 은닉층)로 학습한다. 공정한 비교를 위해 RSMA-PPO와 SDMA-PPO 모두 동일한 학습 상태, posterior "
          "channel 표본 8개, 80 epochs 및 one-step 기대 AoI 목적함수로 사전학습한 뒤, 동일한 PPO curriculum과 "
          "학습 step을 적용한다. 혼합 조건(각도 폭 5/20/60°, SNR 10/15/20 dB, "
          "$n$=128/256)과 목표 조건(20°, 15 dB, $n$=128)에서 각각 524,288 step 학습한다. 목표 조건 학습에서 "
          "에피소드 보상은 5개 seed 모두 증가해 평균 −1033에서 −797로 수렴했다(그림 1). 이는 탐색 중 slot당 "
          "평균 AoI 약 2.02에서 1.56으로의 감소에 해당한다."),
    ("figure", "return"),
    ("gap",),
    ("h", "Ⅳ. 실험 결과"),
    ("p", "평가는 결정적 정책을 6개 사용자 순열에서 768 slot 실행하고 처음 256 slot을 제외하며, 모든 방법이 "
          "같은 채널, CSIT 오차, ACK 난수를 공유한다. CI는 5개 seed 평균의 95% Student-t 구간이다. 기준선은 "
          "(i) $\\alpha=\\phi_k=0$으로 고정하고 $q_k$만 학습하는 SDMA-PPO(동일 데이터·objective-only 사전학습·"
          "posterior 표본 수·epoch·PPO curriculum·학습 step), (ii) 최대 AoI 사용자의 패킷 전체를 common으로 "
          "보내고 $\\alpha$를 검증으로 0.70으로 "
          "정한 고정 규칙 RSMA, (iii) 최대 AoI 사용자 한 명에게 전체 전력을 주는 단일 사용자 MRT(평균 AoI "
          "하한 2)다."),
    ("table", "target"),
    ("p", "동일한 사전학습과 PPO 조건의 RSMA-PPO는 SDMA-PPO보다 평균 AoI가 16.01%(±3.40%p) "
          "낮았고 5개 seed와 6개 순열의 30개 조합 모두에서 더 낮았다. 행동 차원과 최적화 난도가 달라 이를 common "
          "stream만의 인과효과로 단정하지 않는다. 고정 규칙 RSMA와 "
          "단일 사용자 대비 개선은 각각 11.63%(±1.63%p), 26.87%(±1.34%p)였다."),
    ("table", "rank"),
    ("p", "표 2에서 정책은 AoI가 가장 큰 사용자에게 private 전력 비율을 0.871까지 집중했고, common payload는 "
          "AoI 순위에 따라 다르게 배분했다. 최대 AoI가 2에서 3 이상으로 커지면 $\\alpha$는 0.768에서 0.669로 "
          "줄었다. 그림 2와 같이 feasible 조건(SNR≥10 dB, $n$≥128)에서 $\\alpha$는 SNR과 함께 0.47에서 "
          "0.80으로 증가했고, LoS 방향이 더 분리된 60°에서는 0.51로 감소해 SDMA 쪽으로 이동했다."),
    ("figure", "alpha"),
    ("p", "이 feasible 조건의 평균 AoI는 1.739로 단일 사용자(2.000)보다 13.07%(±5.26%p) 낮았다. 반면 PPO "
          "학습에서 제외한 SNR 5 dB 또는 $n$=64 조건에서는 512-slot 측정 평균 AoI가 148.2로 단일 사용자(67.9)보다 나빴다."),
    ("gap",),
    ("h", "Ⅴ. 결 론"),
    ("p", "본 논문은 SDMA 경계를 포함하는 연속 행동으로 FBL RSMA의 전력·payload 분할을 학습하는 PPO 정책을 "
          "제시하였다. 목표 조건에서 제안 정책은 동일한 사전학습과 PPO 조건의 SDMA-PPO보다 16.01%, 두 AoI 규칙 기준선보다 "
          "11.63~26.87% 낮은 평균 AoI를 보였고, AoI 순위, "
          "SNR, 사용자 각도에 따라 행동을 조절했다. 다만 결과는 한 CSIT 오차 수준과 고정 beam 규칙에 한정되며, "
          "저SNR·짧은 blocklength에서의 열세는 학습 분포 밖 조건과 전체 사용자 활성화 설계의 영향을 분리하지 "
          "못했다. 향후 독립 평가 난수와 사용자 선택을 포함한 확장을 수행할 예정이다."),
]

TABLES = {
    "target": {
        "caption": "표 1. 목표 조건(20°, 15 dB, n=128)의 성능",
        "rows": [
            ("방법", "평균 AoI", "95% CI", "p95", "성공확률"),
            ("RSMA-PPO", "1.463", "±0.015", "2.03", "0.577"),
            ("고정 규칙 RSMA", "1.656", "±0.046*", "3.00", "0.502"),
            ("SDMA-PPO", "1.744", "±0.074", "3.17", "0.447"),
            ("단일 사용자 MRT", "2.000", "–", "3.00", "0.333"),
        ],
        "widths": [1650, 850, 800, 550, 800],
        "bold_rows": (0, 1),
        "note": ("*규칙 정책은 6개 순열 간 표준편차. p95는 궤적별 AoI 95백분위수의 평균, 성공확률은 "
                 "사용자·slot별 PHY 성공확률의 평균."),
    },
    "rank": {
        "caption": "표 2. 목표 조건에서 현재 AoI 순위별 평균 행동",
        "rows": [
            ("AoI 순위", "private 전력 $q_k$", "common payload $\\phi_k$"),
            ("가장 낮음", "0.047 (±0.094)", "0.539 (±0.188)"),
            ("중간", "0.082 (±0.074)", "0.303 (±0.139)"),
            ("가장 높음", "0.871 (±0.167)", "0.284 (±0.213)"),
        ],
        "widths": [1050, 1750, 1950],
        "bold_rows": (0,),
        "note": "괄호는 5개 seed의 95% CI. AoI 동률은 사용자 번호 순으로 순위를 매김.",
    },
}

FIGURES = {
    "return": {
        "path": "output/kics/figure_return_curve.png",
        "width_cm": 7.6,
        "caption": ("그림 1. 목표 조건 미세조정의 에피소드(512 slot) 보상. 선은 seed별 8-point 이동평균, "
                    "굵은 선과 음영은 5개 seed 평균 ±1 SD."),
    },
    "alpha": {
        "path": "output/kics/figure_alpha_feasible.png",
        "width_cm": 7.6,
        "caption": ("그림 2. Feasible 조건(18개)의 평균 common 전력 비율 α. (a) SNR별, (b) 사용자 각도 "
                    "폭별, error bar는 95% CI."),
    },
}

ACK_TEXT = None  # set to a sponsor sentence to restore the template's ACKNOWLEDGMENT section

REFERENCES = [
    "[1] S. Kaul, R. Yates, and M. Gruteser, \"Real-time status: How often should one update?,\" "
    "in Proc. IEEE INFOCOM, pp. 2731-2735, 2012.",
    "[2] Y. Mao, B. Clerckx, and V. O. K. Li, \"Rate-splitting multiple access for downlink communication "
    "systems: Bridging, generalizing, and outperforming SDMA and NOMA,\" EURASIP J. Wireless Commun. "
    "Netw., vol. 2018, no. 1, 2018.",
    "[3] Y. Polyanskiy, H. V. Poor, and S. Verdú, \"Channel coding rate in the finite blocklength regime,\" "
    "IEEE Trans. Inf. Theory, vol. 56, no. 5, pp. 2307-2359, 2010.",
    "[4] Y. Xu, Y. Mao, O. Dizdar, and B. Clerckx, \"Rate-splitting multiple access with finite blocklength "
    "for short-packet and low-latency downlink communications,\" IEEE Trans. Veh. Technol., vol. 71, "
    "no. 11, pp. 12333-12337, 2022.",
    "[5] E. Kaya, \"Reinforcement learning based rate splitting for minimizing age of information in finite "
    "blocklength systems,\" M.S. thesis, Middle East Technical Univ., 2025.",
    "[6] J. Schulman, F. Wolski, P. Dhariwal, A. Radford, and O. Klimov, \"Proximal policy optimization "
    "algorithms,\" arXiv:1707.06347, 2017.",
]


# ---------------------------------------------------------------------------
# OMML helpers
# ---------------------------------------------------------------------------
class MathConverter:
    def __init__(self, xsl_path: Path, size_half_points: int = 18) -> None:
        self.xslt = etree.XSLT(etree.parse(str(xsl_path)))
        self.size = size_half_points

    def __call__(self, latex: str) -> etree._Element:
        mathml = latex2mathml.converter.convert(latex)
        omml = self.xslt(etree.fromstring(mathml)).getroot()
        for run in omml.iter(f"{{{M_NS}}}r"):
            rpr = etree.Element(f"{{{W_NS}}}rPr")
            fonts = etree.SubElement(rpr, f"{{{W_NS}}}rFonts")
            fonts.set(f"{{{W_NS}}}ascii", "Cambria Math")
            fonts.set(f"{{{W_NS}}}hAnsi", "Cambria Math")
            for tag in ("sz", "szCs"):
                etree.SubElement(rpr, f"{{{W_NS}}}{tag}").set(f"{{{W_NS}}}val", str(self.size))
            text = run.find(f"{{{M_NS}}}t")
            if text is None:
                run.append(rpr)
            else:
                text.addprevious(rpr)
        return omml


def _m(tag: str, parent=None):
    el = etree.Element(f"{{{M_NS}}}{tag}") if parent is None else etree.SubElement(parent, f"{{{M_NS}}}{tag}")
    return el


def _math_run(parent, text: str, size: int) -> None:
    run = _m("r", parent)
    rpr = etree.SubElement(run, f"{{{W_NS}}}rPr")
    fonts = etree.SubElement(rpr, f"{{{W_NS}}}rFonts")
    fonts.set(f"{{{W_NS}}}ascii", "Cambria Math")
    fonts.set(f"{{{W_NS}}}hAnsi", "Cambria Math")
    for tag in ("sz", "szCs"):
        etree.SubElement(rpr, f"{{{W_NS}}}{tag}").set(f"{{{W_NS}}}val", str(size))
    _m("t", run).text = text


def numbered_display(omath: etree._Element, number: str, size: int = 18) -> etree._Element:
    """Wrap an equation as Word stores ``equation#(n)``: a display equation
    with a right-aligned number (oMathPara > oMath > eqArr with maxDist)."""
    para = _m("oMathPara")
    outer = _m("oMath", para)
    arr = _m("eqArr", outer)
    pr = _m("eqArrPr", arr)
    _m("maxDist", pr).set(f"{{{M_NS}}}val", "1")
    body = _m("e", arr)
    for child in list(omath):
        body.append(child)
    _math_run(body, "#", size)
    delim = _m("d", body)
    _math_run(_m("e", delim), number, size)
    return para


# ---------------------------------------------------------------------------
# document helpers
# ---------------------------------------------------------------------------
def set_text(paragraph, text: str) -> None:
    """Keep the first run's formatting and replace the paragraph text."""
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def clone_paragraph(doc, template_ppr, jc: str | None = None):
    """Append an empty paragraph carrying a copy of a template paragraph's pPr."""
    new = doc.add_paragraph()
    if new._p.pPr is not None:
        new._p.remove(new._p.pPr)
    ppr = copy.deepcopy(template_ppr)
    new._p.insert(0, ppr)
    if jc is not None:
        el = ppr.find(qn("w:jc"))
        if el is None:
            el = etree.Element(qn("w:jc"))
            anchor = ppr.find(qn("w:rPr"))
            anchor.addprevious(el) if anchor is not None else ppr.append(el)
        el.set(qn("w:val"), jc)
    return new


def keep_with_next(paragraph) -> None:
    ppr = paragraph._p.pPr
    if ppr.find(qn("w:keepNext")) is None:
        el = etree.Element(qn("w:keepNext"))
        style = ppr.find(qn("w:pStyle"))
        style.addnext(el) if style is not None else ppr.insert(0, el)


def add_run(paragraph, proto_rpr, text: str, bold: bool = False, size: int | None = None):
    run = paragraph.add_run(text)
    rpr = copy.deepcopy(proto_rpr)
    if run._r.rPr is not None:
        run._r.remove(run._r.rPr)
    run._r.insert(0, rpr)
    if bold:
        run.bold = True
    if size is not None:
        run.font.size = Pt(size)
    return run


def add_rich_text(paragraph, proto_rpr, text: str, math: MathConverter) -> None:
    for i, chunk in enumerate(re.split(r"\$(.+?)\$", text)):
        if not chunk:
            continue
        if i % 2:
            paragraph._p.append(math(chunk))
        else:
            add_run(paragraph, proto_rpr, chunk)


def set_tabs(paragraph, stops: list[tuple[str, int]]) -> None:
    ppr = paragraph._p.pPr
    tabs = etree.Element(qn("w:tabs"))
    for kind, pos in stops:
        tab = etree.SubElement(tabs, qn("w:tab"))
        tab.set(qn("w:val"), kind)
        tab.set(qn("w:pos"), str(pos))
    ppr.find(qn("w:spacing")).addprevious(tabs)


def zero_indent(paragraph) -> None:
    ind = paragraph._p.pPr.find(qn("w:ind"))
    if ind is not None:
        paragraph._p.pPr.remove(ind)


def set_cell_border(cell, **edges) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = etree.SubElement(tcpr, qn("w:tcBorders"))
    for edge, size in edges.items():
        el = etree.SubElement(borders, qn(f"w:{edge}"))
        el.set(qn("w:val"), "single" if size else "nil")
        if size:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), "000000")


def compatibility_mode(doc) -> int:
    for el in doc.settings.element.iter(qn("w:compatSetting")):
        if el.get(qn("w:name")) == "compatibilityMode":
            return int(el.get(qn("w:val")))
    return 0


def replace_font(doc, old: str, east_asian: str, latin: str) -> int:
    """Map the template font to an East Asian font plus a Latin font.

    The template's eastAsia hint is dropped so that ambiguous-width symbols
    (·, ±, °) use the Latin font instead of full-width Hangul-font glyphs.
    """
    count = 0
    for root in (doc.element, doc.styles.element):
        for fonts in root.iter(qn("w:rFonts")):
            if old not in (fonts.get(qn("w:ascii")), fonts.get(qn("w:eastAsia"))):
                continue
            fonts.set(qn("w:eastAsia"), east_asian)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(attr), latin)
            if fonts.get(qn("w:hint")) is not None:
                del fonts.attrib[qn("w:hint")]
            count += 1
    return count


def build(template: Path, xsl: Path, out: Path) -> None:
    doc = docx.Document(str(template))
    mode = compatibility_mode(doc)
    if mode < 15:
        raise SystemExit(
            f"template is in Word compatibility mode {mode}; convert it to mode 15 first "
            "(native equations break on save in older modes)"
        )
    math = MathConverter(xsl)
    table_math = MathConverter(xsl, size_half_points=16)
    ps = doc.paragraphs
    proto = {name: copy.deepcopy(ps[i].runs[0]._r.rPr) for name, i in (
        ("head", T_HEAD), ("body", T_BODY), ("ref", T_REF), ("title", T_TITLE1),
    )}
    templ = {name: copy.deepcopy(ps[i]._p.pPr) for name, i in (
        ("head", T_HEAD), ("body", T_BODY), ("gap", T_GAP),
        ("ref_head", T_REF_HEAD), ("ref", T_REF),
    )}

    # --- title block (single column) ------------------------------------
    set_text(ps[T_TITLE1], TITLE_KO[0])
    set_text(ps[T_TITLE2], TITLE_KO[1])
    set_text(ps[T_AUTH], AUTHORS_KO)
    set_text(ps[T_AFFIL], AFFIL_KO)
    set_text(ps[T_MAIL], EMAIL)
    set_text(ps[T_ETITLE1], TITLE_EN[0])
    set_text(ps[T_ETITLE2], TITLE_EN[1])
    set_text(ps[T_EAUTH], AUTHORS_EN)
    set_text(ps[T_EAFFIL], AFFIL_EN)
    set_text(ps[T_ABS_HEAD], "요 약")
    abstract = ps[T_ABS]
    set_text(abstract, ABSTRACT)
    ppr = abstract._p.pPr
    ppr.find(qn("w:jc")).set(qn("w:val"), "both")
    ind = etree.Element(qn("w:ind"))
    ind.set(qn("w:left"), "567")
    ind.set(qn("w:right"), "567")
    ind.set(qn("w:firstLineChars"), "100")
    ind.set(qn("w:firstLine"), "180")
    ppr.find(qn("w:jc")).addprevious(ind)
    # drop the dotted filler lines and extra blank lines inside the title section
    for i in (18, 17, 16, 15, 13):
        el = ps[i]._p
        el.getparent().remove(el)

    # --- remove the template body (two-column section) -------------------
    for p in doc.paragraphs:
        if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None:
            break_at = p._p
            break
    body = break_at.getparent()
    started = False
    for el in list(body):
        if el is break_at:
            started = True
            continue
        if started and el.tag != qn("w:sectPr"):
            body.remove(el)

    # --- two-column body ---------------------------------------------------
    for block in BODY:
        kind = block[0]
        if kind == "h":
            p = clone_paragraph(doc, templ["head"])
            keep_with_next(p)
            add_run(p, proto["head"], block[1], bold=True)
        elif kind == "gap":
            p = clone_paragraph(doc, templ["gap"])
            spacing = p._p.pPr.find(qn("w:spacing"))
            spacing.set(qn("w:line"), "120")
        elif kind == "p":
            p = clone_paragraph(doc, templ["body"])
            add_rich_text(p, proto["body"], block[1], math)
        elif kind == "eq":
            p = clone_paragraph(doc, templ["body"], jc="center")
            zero_indent(p)
            spacing = p._p.pPr.find(qn("w:spacing"))
            spacing.set(qn("w:before"), "20")
            spacing.set(qn("w:after"), "20")
            p._p.append(numbered_display(math(block[1]), block[2]))
        elif kind == "table":
            add_table(doc, templ, proto, TABLES[block[1]], table_math)
        elif kind == "figure":
            add_figure(doc, templ, proto, FIGURES[block[1]])

    p = clone_paragraph(doc, templ["gap"])
    p._p.pPr.find(qn("w:spacing")).set(qn("w:line"), "120")
    if ACK_TEXT:
        p = clone_paragraph(doc, templ["ref_head"])
        add_run(p, proto["head"], "ACKNOWLEDGMENT", bold=True)
        p = clone_paragraph(doc, templ["body"])
        add_run(p, proto["body"], ACK_TEXT)
        p = clone_paragraph(doc, templ["gap"])
        p._p.pPr.find(qn("w:spacing")).set(qn("w:line"), "120")
    p = clone_paragraph(doc, templ["ref_head"])
    keep_with_next(p)
    add_run(p, proto["head"], "참 고 문 헌", bold=True)
    for ref in REFERENCES:
        p = clone_paragraph(doc, templ["ref"])
        p._p.pPr.find(qn("w:spacing")).set(qn("w:after"), "20")
        add_run(p, proto["ref"], ref)
    replace_font(doc, TEMPLATE_FONT, BODY_FONT, LATIN_FONT)
    doc.save(str(out))


def caption(doc, templ, proto, text: str, before: int, after: int) -> None:
    p = clone_paragraph(doc, templ["body"], jc="center")
    zero_indent(p)
    spacing = p._p.pPr.find(qn("w:spacing"))
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    keep_with_next(p)
    label, rest = text.split(". ", 1)
    add_run(p, proto["body"], label + ". ", bold=True, size=8.5)
    add_run(p, proto["body"], rest, size=8.5)


def set_cell_margins(table, twips: int) -> None:
    tblpr = table._tbl.tblPr
    mar = etree.SubElement(tblpr, qn("w:tblCellMar"))
    for edge in ("left", "right"):
        el = etree.SubElement(mar, qn(f"w:{edge}"))
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")


def add_table(doc, templ, proto, spec: dict, math: MathConverter) -> None:
    rows = spec["rows"]
    caption(doc, templ, proto, spec["caption"], before=80, after=40)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_margins(table, 40)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Cm(spec["widths"][c] / 567)
            p = cell.paragraphs[0]
            p.alignment = 0 if c == 0 else 1
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(0.5)
            for i, chunk in enumerate(re.split(r"\$(.+?)\$", value)):
                if i % 2:
                    p._p.append(math(chunk))
                elif chunk:
                    add_run(p, proto["body"], chunk, bold=r in spec["bold_rows"], size=8)
            top = 12 if r == 0 else 0
            bottom = 12 if r == len(rows) - 1 else (6 if r == 0 else 0)
            set_cell_border(cell, top=top, bottom=bottom, left=0, right=0)
    note = clone_paragraph(doc, templ["body"])
    zero_indent(note)
    note._p.pPr.find(qn("w:spacing")).set(qn("w:after"), "80")
    add_run(note, proto["body"], spec["note"], size=7.5)


def add_figure(doc, templ, proto, spec: dict) -> None:
    p = clone_paragraph(doc, templ["body"], jc="center")
    zero_indent(p)
    p._p.pPr.find(qn("w:spacing")).set(qn("w:before"), "80")
    keep_with_next(p)
    p.add_run().add_picture(str(ROOT / spec["path"]), width=Cm(spec["width_cm"]))
    caption(doc, templ, proto, spec["caption"], before=20, after=80)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--template", default=str(ROOT / "output/kics/kics_template_word2013.docx"),
                    help="KICS sample converted to .docx in compatibility mode 15")
    ap.add_argument("--xsl", required=True, help="Office MML2OMML.XSL")
    ap.add_argument("--out", default=str(ROOT / "output/kics/RSMA_PPO_AoI_KICS_2p.docx"))
    args = ap.parse_args()
    build(Path(args.template), Path(args.xsl), Path(args.out))
    print("saved", args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
